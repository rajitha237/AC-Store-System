from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from app.services.documents.financial_report import (
    FinancialReportPDFData,
)


def build_financial_report_excel(
    data: FinancialReportPDFData,
) -> bytes:
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Financial Summary"

    currency_format = (
        f'"{data.currency_code}" '
        '#,##0.00;[Red]-'
        f'"{data.currency_code}" '
        '#,##0.00'
    )

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D",
    )

    section_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="F2F4F7",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    bold_font = Font(
        bold=True,
    )

    thin_border = Border(
        left=Side(
            style="thin",
            color="D0D5DD",
        ),
        right=Side(
            style="thin",
            color="D0D5DD",
        ),
        top=Side(
            style="thin",
            color="D0D5DD",
        ),
        bottom=Side(
            style="thin",
            color="D0D5DD",
        ),
    )

    sheet.merge_cells(
        "A1:B1"
    )

    sheet["A1"] = (
        data.company_name
    )

    sheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )

    sheet["A1"].fill = (
        title_fill
    )

    sheet["A1"].alignment = (
        Alignment(
            horizontal="center",
            vertical="center",
        )
    )

    sheet.row_dimensions[1].height = 26

    sheet.merge_cells(
        "A2:B2"
    )

    sheet["A2"] = (
        "FINANCIAL REPORT"
    )

    sheet["A2"].font = Font(
        bold=True,
        size=13,
    )

    sheet["A2"].alignment = (
        Alignment(
            horizontal="center",
        )
    )

    sheet.merge_cells(
        "A3:B3"
    )

    sheet["A3"] = (
        f"Period: "
        f"{data.date_from.isoformat()} "
        f"to "
        f"{data.date_to.isoformat()}"
    )

    sheet["A3"].alignment = (
        Alignment(
            horizontal="center",
        )
    )

    sheet.merge_cells(
        "A4:B4"
    )

    sheet["A4"] = (
        f"Timezone: "
        f"{data.timezone_name}"
    )

    sheet["A4"].alignment = (
        Alignment(
            horizontal="center",
        )
    )

    row = 6

    def add_section(
        title: str,
        rows: list[
            tuple[
                str,
                object,
                bool,
            ]
        ],
    ) -> None:
        nonlocal row

        sheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=2,
        )

        section_cell = (
            sheet.cell(
                row=row,
                column=1,
            )
        )

        section_cell.value = title
        section_cell.font = (
            bold_font
        )
        section_cell.fill = (
            section_fill
        )

        row += 1

        label_header = (
            sheet.cell(
                row=row,
                column=1,
                value="Account",
            )
        )

        amount_header = (
            sheet.cell(
                row=row,
                column=2,
                value=(
                    f"Amount "
                    f"({data.currency_code})"
                ),
            )
        )

        for cell in (
            label_header,
            amount_header,
        ):
            cell.font = bold_font
            cell.fill = header_fill
            cell.border = thin_border

        amount_header.alignment = (
            Alignment(
                horizontal="right",
            )
        )

        row += 1

        for (
            label,
            amount,
            strong,
        ) in rows:
            label_cell = (
                sheet.cell(
                    row=row,
                    column=1,
                    value=label,
                )
            )

            amount_cell = (
                sheet.cell(
                    row=row,
                    column=2,
                    value=float(amount),
                )
            )

            label_cell.border = (
                thin_border
            )

            amount_cell.border = (
                thin_border
            )

            amount_cell.number_format = (
                currency_format
            )

            amount_cell.alignment = (
                Alignment(
                    horizontal="right",
                )
            )

            if strong:
                label_cell.font = (
                    bold_font
                )

                amount_cell.font = (
                    bold_font
                )

                label_cell.fill = (
                    total_fill
                )

                amount_cell.fill = (
                    total_fill
                )

            row += 1

        row += 1

    add_section(
        "Profit & Loss",
        [
            (
                "Gross sales",
                data.gross_sales,
                False,
            ),
            (
                "Sales credits / reductions",
                data.sales_credits,
                False,
            ),
            (
                "Net revenue",
                data.net_revenue,
                True,
            ),
            (
                "Cost of goods sold",
                data.cost_of_goods_sold,
                False,
            ),
            (
                "Gross profit",
                data.gross_profit,
                True,
            ),
            (
                "Recorded operating expenses",
                data.operating_expenses,
                False,
            ),
            (
                "Net profit",
                data.net_profit,
                True,
            ),
        ],
    )

    add_section(
        "Cash Flow",
        [
            (
                "Customer collections",
                data.customer_collections,
                False,
            ),
            (
                "Manual cash in",
                data.manual_cash_in,
                False,
            ),
            (
                "Total cash in",
                data.total_cash_in,
                True,
            ),
            (
                "Supplier payments",
                data.supplier_payments,
                False,
            ),
            (
                "Manual cash out",
                data.manual_cash_out,
                False,
            ),
            (
                "Total cash out",
                data.total_cash_out,
                True,
            ),
            (
                "Net cash flow",
                data.net_cash_flow,
                True,
            ),
        ],
    )

    add_section(
        "Purchasing & Outstanding",
        [
            (
                "Supplier invoices",
                data.supplier_invoices,
                False,
            ),
            (
                "Customer receivables",
                data.customer_receivables,
                True,
            ),
            (
                "Supplier payables",
                data.supplier_payables,
                True,
            ),
        ],
    )

    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    note = sheet.cell(
        row=row,
        column=1,
    )

    note.value = (
        "Accounting note: "
        "Profit is reported separately "
        "from cash collections. "
        "Customer collections may include "
        "payments from earlier periods. "
        "Operating expenses include only "
        "recorded active manual cash-out "
        "entries."
    )

    note.alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    note.font = Font(
        italic=True,
        color="667085",
    )

    sheet.column_dimensions[
        get_column_letter(1)
    ].width = 42

    sheet.column_dimensions[
        get_column_letter(2)
    ].width = 24

    sheet.freeze_panes = "A7"

    sheet.sheet_view.showGridLines = (
        False
    )

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    return buffer.getvalue()
